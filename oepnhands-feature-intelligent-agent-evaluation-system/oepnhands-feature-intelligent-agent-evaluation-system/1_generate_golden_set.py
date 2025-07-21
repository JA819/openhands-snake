import requests
import json
from typing import List, Dict
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import random

class DocxToQADocx:
    def __init__(self, api_key: str):
        self.api_key = api_key
        self.chat_url = "https://api.siliconflow.cn/v1/chat/completions"  # 请确保链接正确
        self.headers = {
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json"
        }
    
    def read_docx(self, filepath: str) -> List[str]:
        """读取并返回文档中的所有段落"""
        from docx import Document
        doc = Document(filepath)
        paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
        return paragraphs

    def generate_qa_pairs(self, text: str) -> List[Dict]:
        """调用API生成问答对"""
        prompt = f"""请基于以下文本生成专业问答对，要求：
        1. 问题要覆盖文本核心内容
        2. 答案必须直接来自文本，并且涵盖完整内容
        3. 使用中文回答
        4. 格式为JSON，包含question和answer字段
        5. 禁止生成与文本内容无关的问题
        6. 禁止编造答案
        7. 禁止将提示词写入问答对文档
        8.减少含有“是否”一类问题的提问
        文本内容：
        {text}

        返回格式示例：
        {{
            "questions": [
                {{
                    "question": "问题内容",
                    "answer": "答案内容"
                }}
            ]
        }}"""
        
        payload = {
            "model": "Qwen/Qwen3-32B",
            "messages": [
                {"role": "system", "content": "你是专业文档处理助手，擅长从技术文本中提取关键信息生成问答对"},
                {"role": "user", "content": prompt}
            ],
            "temperature": 0.3,
            "response_format": {"type": "json_object"}
        }
        
        try:
            response = requests.post(self.chat_url, json=payload, headers=self.headers, timeout=30)
            response.raise_for_status()
            result = response.json()
            content = result.get('choices', [{}])[0].get('message', {}).get('content', '{}')
            qa_data = json.loads(content)
            return qa_data.get('questions', [])
        except json.JSONDecodeError:
            print("API返回的JSON格式不正确")
            return []
        except Exception as e:
            print(f"生成问答对失败: {str(e)}")
            print(f"API响应: {response.text if 'response' in locals() else ''}")
            return []

    def create_qa_excel(self, qa_pairs: List[Dict], output_path: str):
        """生成规范的问答对Excel文档"""
        wb = Workbook()
        ws = wb.active
        ws.title = "问答对"
        
        # 设置表头
        ws.append(["问题", "答案"])
        for cell in ws[1]:
            cell.font = Font(bold=True, size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 添加问答内容
        for qa in qa_pairs:
            question = qa.get('question', '无问题内容')
            answer = qa.get('answer', '无答案内容')
            ws.append([question, answer])
        
        # 自动调整列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter  # 获取列字母
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width
        
        wb.save(output_path)
        print(f"已生成问答Excel: {os.path.abspath(output_path)}")

    def process_document(self, input_path: str, output_path: str):
        """完整处理流程"""
        print(f"正在处理文档: {input_path}")
        paragraphs = self.read_docx(input_path)
        
        if not paragraphs:
            print("警告：文档中没有有效段落")
            return []
        
        # 随机选择20个段落
        random_paragraphs = random.sample(paragraphs, min(20, len(paragraphs)))
        
        print("正在生成问答对...")
        all_qa = []
        for i, para in enumerate(random_paragraphs, 1):
            qa_pairs = self.generate_qa_pairs(para)
            for qa in qa_pairs:
                all_qa.append(qa)
            print(f"进度: {i}/{len(random_paragraphs)} 段落，当前段落生成 {len(qa_pairs)} 个问答对")
        
        if not all_qa:
            print("警告：未生成任何问答对，请检查API调用或文档内容")
            return []
        
        self.create_qa_excel(all_qa, output_path)
        return all_qa

# 使用示例
if __name__ == "__main__":
    API_KEY = "sk-elnmwevbokezmyjvyafilsfsvdgwqbgsrjvlrnfhzsodtakc"  # 替换为实际API密钥
    processor = DocxToQADocx(API_KEY)
    
    input_docx = "C:/Users/19541/Desktop/Internship services/Internship services/Dify/learnproject/learn HTTP/延期换证办事指南.docx"# 输入文档路径
    output_excel = "C:/Users/19541/Desktop/Internship services/Internship services/Dify/learnproject/learn HTTP/问答对1.xlsx" # 输出Excel文件路径
    
    qa_pairs = processor.process_document(input_docx, output_excel)
    if qa_pairs:
        print(f"处理完成，共生成 {len(qa_pairs)} 个问答对")
    else:
        print("处理完成，但未生成有效问答对")
